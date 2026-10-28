"""Convierte Libro1.xlsx (Control de inventario Internet K-ROL) en CSV listos para importar.
Uso (desde la raíz del repo): python "plan de implementacion/migrar_excel.py" Libro1.xlsx
Genera en "plan de implementacion/": productos_importar.csv, activos_fijos_importar.csv, avisos_revision.csv
"""
import csv, re, sys
from openpyxl import load_workbook

src = sys.argv[1] if len(sys.argv) > 1 else "Libro1.xlsx"
ws = load_workbook(src).active

limpiar = lambda t: re.sub(r"\s+", " ", str(t)).strip()
productos, activos, avisos, vistos = [], [], [], set()
seccion, categoria = "mercaderia", None

for fila, (a, b, c, d, e) in enumerate(ws.iter_rows(min_row=5, max_col=5, values_only=True), start=5):
    if not b:
        continue
    nombre = limpiar(b)
    if "NO CORRIENTE" in nombre.upper():
        seccion, categoria = "activos", None
        continue
    if a is None and c is None and d is None:          # fila de encabezado/categoría
        if nombre.upper() != "MERCADERÍAS":
            categoria = nombre.capitalize()
        continue
    if seccion == "mercaderia":
        stock = a if isinstance(a, (int, float)) else 0
        clave = (categoria, nombre.lower())
        if clave in vistos:
            avisos.append((fila, nombre, "Nombre repetido en la misma categoría"))
        vistos.add(clave)
        if c is None or d is None:
            avisos.append((fila, nombre, "Sin costo o sin precio de venta"))
        elif d <= c:
            avisos.append((fila, nombre, f"Precio de venta ({d}) menor o igual al costo ({c})"))
        if stock > 1000:
            avisos.append((fila, nombre, f"Cantidad muy alta ({stock}); verificar posible error de digitación"))
        productos.append([f"P{len(productos)+1:04d}", categoria, nombre, c, d, stock, fila])
    else:
        activos.append([categoria or "Mobiliario y Equipo", nombre, a, c, fila])

def guardar(nombre, cabecera, filas):
    with open(nombre, "w", newline="", encoding="utf-8-sig") as f:
        w = csv.writer(f); w.writerow(cabecera); w.writerows(filas)

d = "plan de implementacion/"
guardar(d+"productos_importar.csv", ["codigo","categoria","nombre","costo","precio","stock_inicial","fila_excel"], productos)
guardar(d+"activos_fijos_importar.csv", ["grupo","nombre","cantidad","costo_unitario","fila_excel"], activos)
guardar(d+"avisos_revision.csv", ["fila_excel","producto","aviso"], avisos)
print(len(productos), "productos |", len({p[1] for p in productos}), "categorías |", len(activos), "activos |", len(avisos), "avisos")
